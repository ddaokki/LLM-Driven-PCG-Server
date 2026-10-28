"""
scripts/build-docx.py <in.md> <out.docx>
test-report.js가 만든 결과서 markdown(제목/표/문단/인용/목록만 사용)을 python-docx로 Word 문서로 변환한다.
제출용 문서 양식(테스트 명세서.docx)과 같은 순서(개요 → 환경 → 결과 → 세부 → 비고)를 유지한다.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_inline(paragraph, text):
    """`code`와 **bold**만 처리하는 최소 인라인 렌더러."""
    pos = 0
    for m in re.finditer(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def flush_table(doc, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        set_cell_bg(cell, "D9E2F3")
    for row in body:
        cells = table.add_row().cells
        for i in range(len(header)):
            val = row[i] if i < len(row) else ""
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_inline(p, val)
            for r in p.runs:
                r.font.size = Pt(9)
                if val.startswith("FAIL"):
                    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                    r.bold = True
                elif val == "PASS":
                    r.font.color.rgb = RGBColor(0x1E, 0x7B, 0x34)
    doc.add_paragraph()


def main(src, dst):
    lines = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10)

    table_rows = []
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            if re.match(r"^\|\s*-+", line):
                continue
            cells = [c.strip().replace("\\|", "|") for c in line.strip("|").split("|")]
            table_rows.append(cells)
            continue
        flush_table(doc, table_rows)
        table_rows = []
        if not line:
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:])
            r.italic = True
            r.font.size = Pt(9)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
    flush_table(doc, table_rows)
    doc.save(dst)
    print(f"saved {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
